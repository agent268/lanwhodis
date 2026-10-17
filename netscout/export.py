"""Result writers: json, html, md, csv, xlsx, pdf, txt. Shared by the CLI and GUI."""
from __future__ import annotations

import csv
import html
import io
import json
import socket
from datetime import datetime
from pathlib import Path
from typing import Iterable

from . import icons
from .model import Device

FORMATS = ("json", "html", "md", "csv", "xlsx", "pdf", "txt")
BINARY = {"xlsx", "pdf"}
EXT_TO_FORMAT = {".json": "json", ".html": "html", ".htm": "html", ".md": "md", ".markdown": "md",
                 ".csv": "csv", ".xlsx": "xlsx", ".pdf": "pdf", ".txt": "txt"}

COLUMNS = [("IP", "ip"), ("Name", "name"), ("Type", "type_label"), ("OS", "os"), ("Vendor", "vendor"), ("MAC", "mac"),
           ("Model", "model"), ("Confidence", "conf"), ("Open ports", "ports"), ("Services", "services"), ("Web UI", "web_url"), ("Notes", "notes")]


def format_for_path(path: str | Path) -> str | None:
    return EXT_TO_FORMAT.get(Path(path).suffix.lower())


def _rows(devices: Iterable[Device]) -> list[dict]:
    rows = []
    for d in devices:
        x = d.to_dict()
        x["conf"] = f"{int(round(d.confidence * 100))}%"
        x["ports"] = ", ".join(map(str, x["open_ports"]))
        x["services"] = ", ".join(s.split(":", 1)[-1] for s in x["services"] if not s.startswith("local:"))
        x["notes"] = (x.get("user") or {}).get("notes") or ""
        for k in ("vendor", "mac", "model", "os", "web_url"):
            x[k] = x[k] or ""
        rows.append(x)
    return rows


def _meta(devices: list[Device], meta: dict | None) -> dict:
    m = {"generated": datetime.now().astimezone().isoformat(timespec="seconds"),
         "host": socket.gethostname(), "tool": "netscout", "count": len(devices)}
    m.update(meta or {})
    return m


# ---- text-ish formats -------------------------------------------------------------------

def _json(devices, meta) -> str:
    return json.dumps({"meta": meta, "devices": [d.to_dict() for d in devices]}, indent=2) + "\n"


def _csv(rows, meta) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow([c for c, _ in COLUMNS])
    for r in rows:
        w.writerow([r[k] for _, k in COLUMNS])
    return buf.getvalue()


PROTOCOL_ORDER = ("mdns", "upnp", "ssdp", "wsd", "smb", "snmp", "netbios", "tcp")
PROTOCOL_LABEL = {"mdns": "mDNS / Bonjour", "upnp": "UPnP", "ssdp": "SSDP", "wsd": "WS-Discovery",
                  "smb": "SMB / Windows networking", "snmp": "SNMP", "netbios": "NetBIOS",
                  "tcp": "Open TCP ports"}


def grouped_services(r: dict) -> list[tuple[str, list[dict]]]:
    """Advertised services of one device row, grouped by protocol in a stable, readable order."""
    out = []
    for proto in PROTOCOL_ORDER:
        items = [s for s in r["advertised"] if s["protocol"] == proto]
        if items:
            out.append((PROTOCOL_LABEL[proto], items))
    return out


def service_line(s: dict) -> str:
    bits = [s["label"]]
    if s["name"]:
        bits.append(f'"{s["name"]}"')
    if s["port"] and s["protocol"] != "tcp":
        bits.append(f':{s["port"]}')
    if s["url"]:
        bits.append(s["url"] + (" [SOAP]" if s["soap"] else ""))
    return " ".join(bits)


def _service_blocks(rows, md: bool) -> list[str]:
    out = []
    for r in rows:
        groups = grouped_services(r)
        if not groups:
            continue
        out += ["", f"## {r['ip']} - {r['name']}" if md else f"{r['ip']}  {r['name']}  ({r['type_label']})"]
        for label, items in groups:
            out.append(f"**{label}**" if md else f"  {label}")
            out += [(f"- {service_line(s)}" if md else f"    {service_line(s)}") for s in items]
    return out


def _md_cell(s) -> str:
    return str(s).replace("|", "\\|").replace("\n", " ")


def _md(rows, meta) -> str:
    out = [f"# Network scan — {meta['host']}", "",
           f"Generated {meta['generated']} · {meta['count']} devices"
           + (f" · {meta['networks']}" if meta.get("networks") else ""), "",
           "| " + " | ".join(c for c, _ in COLUMNS) + " |",
           "|" + "|".join("---" for _ in COLUMNS) + "|"]
    for r in rows:
        out.append("| " + " | ".join(_md_cell(r[k]) for _, k in COLUMNS) + " |")
    if meta.get("with_services"):
        out += _service_blocks(rows, md=True)
    return "\n".join(out) + "\n"


def _txt(rows, meta) -> str:
    cols = [("IP", "ip", 15), ("Name", "name", 24), ("Type", "type_label", 17), ("OS", "os", 22),
            ("Vendor", "vendor", 22), ("MAC", "mac", 17), ("Conf", "conf", 5)]
    fmt = lambda vals: "  ".join(str(v)[:w].ljust(w) for v, (_, _, w) in zip(vals, cols)).rstrip()
    out = [f"Network scan — {meta['host']} — {meta['generated']} — {meta['count']} devices", "",
           fmt([c for c, _, _ in cols]), fmt(["-" * w for _, _, w in cols])]
    out += [fmt([r[k] for _, k, _ in cols]) for r in rows]
    if meta.get("with_services"):
        out += _service_blocks(rows, md=False)
    return "\n".join(out) + "\n"


def _html(rows, meta) -> str:
    e = html.escape
    cards = []
    for r in rows:
        ev = "".join(f"<li><b>{e(x['source'])}</b> → {e(x['device_type'])} "
                     f"({int(x['weight'] * 100)}%) <span>{e(str(x['detail']))}</span></li>"
                     for x in r["evidence"][:8])
        sv = "".join(f"<li><b>{e(label)}</b><ul>" + "".join(f"<li>{e(service_line(s))}</li>" for s in items)
                     + "</ul></li>" for label, items in grouped_services(r))
        cards.append(
            f'<tr data-q="{e((r["ip"] + r["name"] + r["type_label"] + r["vendor"] + r["mac"]).lower())}">'
            f'<td class="ic">{icons.svg(r["type"], 32)}</td>'
            f'<td><b>{e(r["name"])}</b><br><small>{e(r["ip"])}</small></td>'
            f'<td>{e(r["type_label"])}<br><small>'
            + (f'edited · auto: {e(r["user"]["auto_type_label"])}' if r.get("user") and r["user"]["type"]
               else e(r["conf"]))
            + (f'<br>{e(r["notes"])}' if r["notes"] else "") + '</small></td>'
            f'<td>{e(r["os"])}</td><td>{e(r["vendor"])}<br><small>{e(r["mac"])}</small></td>'
            f'<td>{e(r["model"])}'
            + (f'<br><a href="{e(r["web_url"])}" target="_blank" rel="noopener">Web UI</a>' if r["web_url"] else "")
            + f'</td><td><small>{e(r["ports"])}</small></td>'
            f'<td><details><summary>{len(r["evidence"])} signals</summary><ul>{ev}</ul></details>'
            + (f'<details><summary>{len(r["advertised"])} services</summary><ul>{sv}</ul></details>' if sv else "")
            + '</td></tr>')
    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Network scan — {e(meta['host'])}</title>
<style>
:root{{color-scheme:light dark;--bg:#fff;--fg:#1c2330;--mut:#667;--line:#dde2ea;--card:#f6f8fb}}
@media(prefers-color-scheme:dark){{:root{{--bg:#12161d;--fg:#e6e9ef;--mut:#9aa3b2;--line:#2a313d;--card:#191f29}}}}
body{{margin:0;padding:1.5rem 1rem;background:var(--bg);color:var(--fg);font:14px system-ui,sans-serif}}
main{{max-width:1100px;margin:auto}} h1{{margin:0 0 .25rem}} .sub{{color:var(--mut);margin-bottom:1rem}}
input{{width:100%;max-width:22rem;padding:.5rem .7rem;margin-bottom:1rem;border:1px solid var(--line);
border-radius:8px;background:var(--card);color:var(--fg)}}
.wrap{{overflow-x:auto}} table{{border-collapse:collapse;width:100%;min-width:760px}}
th,td{{text-align:left;padding:.5rem .6rem;border-bottom:1px solid var(--line);vertical-align:top}}
th{{color:var(--mut);font-weight:600;font-size:12px;text-transform:uppercase;letter-spacing:.04em}}
small,summary,li span{{color:var(--mut)}} td.ic{{width:40px}} ul{{margin:.4rem 0;padding-left:1.1rem}}
</style></head><body><main>
<h1>Network scan</h1>
<div class="sub">{e(meta['host'])} · {e(meta['generated'])} · {meta['count']} devices{(' · ' + e(meta['networks'])) if meta.get('networks') else ''}</div>
<input id="q" placeholder="Filter devices…" oninput="for(const r of document.querySelectorAll('tbody tr'))r.hidden=!r.dataset.q.includes(this.value.toLowerCase())">
<div class="wrap"><table><thead><tr><th></th><th>Device</th><th>Type</th><th>OS</th><th>Vendor / MAC</th><th>Model</th><th>Ports</th><th>Evidence</th></tr></thead>
<tbody>{''.join(cards)}</tbody></table></div></main></body></html>
"""


# ---- binary formats ---------------------------------------------------------------------

def _xlsx(rows, meta, devices) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Devices"
    ws.append([c for c, _ in COLUMNS])
    for r in rows:
        ws.append([r[k] for _, k in COLUMNS])
    head = PatternFill("solid", fgColor="1F3A5F")
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head
    for i, (_, k) in enumerate(COLUMNS, 1):
        width = max([len(str(r[k])) for r in rows] + [len(COLUMNS[i - 1][0])])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 48)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=False)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    ev = wb.create_sheet("Evidence")
    ev.append(["IP", "Name", "Source", "Votes for", "Weight", "Detail"])
    for d in devices:
        for x in sorted(d.evidence, key=lambda x: -x.weight):
            ev.append([d.ip, d.display_name, x.source, x.device_type, round(x.weight, 2), x.detail])
    for c in ev[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head
    ev.freeze_panes = "A2"
    ev.auto_filter.ref = ev.dimensions

    sv = wb.create_sheet("Services")
    sv.append(["IP", "Device", "Protocol", "Service", "Name / instance", "Port", "Endpoint", "SOAP", "Details"])
    for r in rows:
        for label, items in grouped_services(r):
            for s in items:
                sv.append([r["ip"], r["name"], label, s["label"], s["name"], s["port"], s["url"] or "",
                           "yes" if s["soap"] else "", "; ".join(f"{k}={v}" for k, v in s["details"].items())[:500]])
    for c in sv[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head
    sv.freeze_panes = "A2"
    sv.auto_filter.ref = sv.dimensions

    info = wb.create_sheet("Scan")
    for k, v in meta.items():
        info.append([k, str(v)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pdf(rows, meta) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    ss = getSampleStyleSheet()
    cell = ss["BodyText"].clone("cell", fontSize=7.5, leading=9)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(letter), leftMargin=.5 * inch, rightMargin=.5 * inch,
                            topMargin=.5 * inch, bottomMargin=.5 * inch, title="Network scan")
    cols = [("IP", "ip"), ("Name", "name"), ("Type", "type_label"), ("OS", "os"), ("Vendor", "vendor"), ("MAC", "mac"),
            ("Model", "model"), ("Conf", "conf"), ("Ports", "ports")]
    data = [[Paragraph(f"<b>{h}</b>", cell) for h, _ in cols]]
    for r in rows:
        data.append([Paragraph(html.escape(str(r[k])), cell) for _, k in cols])
    widths = [.9, 1.45, 1.0, 1.25, 1.35, 1.1, 1.0, .4, .55]
    t = Table(data, colWidths=[w * inch for w in widths], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F5")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7FA")]),
        ("GRID", (0, 0), (-1, -1), .25, colors.HexColor("#C5CCD8")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    sub = f"{meta['host']} · {meta['generated']} · {meta['count']} devices"
    if meta.get("networks"):
        sub += f" · {meta['networks']}"
    doc.build([Paragraph("Network scan", ss["Title"]), Paragraph(html.escape(sub), ss["Normal"]),
               Spacer(1, 10), t])
    return buf.getvalue()


# ---- public API -------------------------------------------------------------------------

def render(devices: list[Device], fmt: str, meta: dict | None = None) -> str | bytes:
    fmt = fmt.lower()
    if fmt not in FORMATS:
        raise ValueError(f"unknown format {fmt!r}; choose from {', '.join(FORMATS)}")
    m = _meta(devices, meta)
    rows = _rows(devices)
    return {
        "json": lambda: _json(devices, m), "csv": lambda: _csv(rows, m), "md": lambda: _md(rows, m),
        "txt": lambda: _txt(rows, m), "html": lambda: _html(rows, m),
        "xlsx": lambda: _xlsx(rows, m, devices), "pdf": lambda: _pdf(rows, m),
    }[fmt]()


def write(devices: list[Device], fmt: str, path: str | Path, meta: dict | None = None) -> Path:
    data = render(devices, fmt, meta)
    p = Path(path)
    if isinstance(data, bytes):
        p.write_bytes(data)
    else:
        p.write_text(data, encoding="utf-8", newline="\n")
    return p
